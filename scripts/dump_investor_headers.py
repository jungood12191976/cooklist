#!/usr/bin/env python3
"""JPX投資部門別売買状況(週次xlsx)の実物の見出し行と、プライムの金額行の全列を data/investor_flows_headers.txt に書き出す(列の配置を実物で確認するための一時スクリプト)。"""
import io, re, requests
from openpyxl import load_workbook
BASE = "https://www.jpx.co.jp"
IDX = BASE + "/markets/statistics-equities/investor-type/index.html"
UA = {"User-Agent": "Mozilla/5.0 (king-the-quest personal app)"}
html = requests.get(IDX, headers=UA, timeout=60).text
links = sorted(set(re.findall(r'href="([^"]*stock_1_w_(\d{8})_(\d{8})\.xlsx)"', html)), key=lambda x: x[1])
out = [f"links: {[(s, e) for _, s, e in links]}"]
href, s, e = links[-1]
ws = load_workbook(io.BytesIO(requests.get(BASE + href, headers=UA, timeout=60).content), data_only=True).worksheets[0]
out.append(f"file {href} / {ws.title} / max_row={ws.max_row} max_col={ws.max_column}")
for r in range(1, 8):
    for c in ws[r]:
        if c.value not in (None, ""):
            out.append(f"hdr {c.coordinate}: {str(c.value)!r}")
mk = None
for r in range(8, ws.max_row + 1):
    b = ws[f"B{r}"].value
    if b: mk = str(b)
    cv = str(ws[f"C{r}"].value or "")
    if mk and "プライム" in mk and cv:
        out.append(f"row {r} market={mk!r} C={cv!r}")
        for c in ws[r]:
            if c.value not in (None, "") and c.column > 3:
                out.append(f"   {c.coordinate}={c.value!r}")
        if "金額" in cv: break
open("data/investor_flows_headers.txt", "w", encoding="utf-8").write("\n".join(out) + "\n")
print("\n".join(out)[:3000])
