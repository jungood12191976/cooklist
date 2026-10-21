#!/usr/bin/env python3
"""JPX 投資部門別売買状況(週次・2026/9/29掲載分以降のxlsx形式)から、海外投資家・個人・信託銀行の
東証3市場(プライム+スタンダード+グロース)・現物の差引(買い−売り)を data/investor_flows.json に追記する。

列(2026-10-21にxlsxの実物で確認。最新週プライムの値はJPX公式PDFと一致):
  海外投資家 = V(法人 差引) + Z(個人 差引)   個人 = N(委託・現金 差引) + R(委託・信用 差引)   信託銀行 = BB(差引)
単位は千円 → 億円(÷100,000)に四捨五入。既に入っている最新週より新しい週だけ追記する(古い週は書き換えない)。
"""
import io, json, re, sys, datetime
import requests
from openpyxl import load_workbook

BASE = "https://www.jpx.co.jp"
IDX = BASE + "/markets/statistics-equities/investor-type/index.html"
UA = {"User-Agent": "Mozilla/5.0 (king-the-quest personal app)"}
JSON_PATH = "data/investor_flows.json"
CHECK_PATH = "data/investor_flows_check.txt"
MARKETS = ("プライム", "スタンダード", "グロース")
COLS = {"foreign": ("V", "Z"), "individual": ("N", "R"), "trustbank": ("BB",)}


def col(ws, letter, row):
    v = ws[f"{letter}{row}"].value
    if not isinstance(v, (int, float)):
        raise ValueError(f"{letter}{row} が数値ではありません: {v!r}")
    return v


def parse(content):
    ws = load_workbook(io.BytesIO(content), data_only=True).worksheets[0]
    # ヘッダーが想定どおりか確認(違ったら何も書かない)
    for letter, word in (("V", "差引"), ("Z", "差引"), ("N", "差引"), ("R", "差引"), ("BB", "差引")):
        if word not in str(ws[f"{letter}7"].value):
            raise ValueError(f"列{letter}の見出しが想定と違います: {ws[f'{letter}7'].value!r}")
    if "海外投資家" not in str(ws["T4"].value) or "信託銀行" not in str(ws["AZ6"].value) or "個人" not in str(ws["L4"].value):
        raise ValueError("見出しの位置が想定と違います")
    tot = {k: 0 for k in COLS}
    seen = []
    market = None
    for r in range(8, ws.max_row + 1):
        b = ws[f"B{r}"].value
        if b:
            market = str(b)
        if "金額" in str(ws[f"C{r}"].value or "") and market and any(m in market for m in MARKETS):
            seen.append(market.split("\n")[0])
            for k, letters in COLS.items():
                tot[k] += sum(col(ws, L, r) for L in letters)
    if len(seen) != 3:
        raise ValueError(f"3市場の金額行が見つかりません: {seen}")
    return {k: int(round(v / 100000)) for k, v in tot.items()}, seen


def monday(d):
    return d - datetime.timedelta(days=d.weekday())


def main():
    doc = json.load(open(JSON_PATH, encoding="utf-8"))
    last_from = max(w["from"] for w in doc["weeks"])
    html = requests.get(IDX, headers=UA, timeout=60).text
    links = re.findall(r'href="([^"]*stock_1_w_(\d{8})_(\d{8})\.xlsx)"', html)
    check = [f"実行 {datetime.datetime.utcnow().isoformat()}Z / 一覧のxlsx {len(links)}件 / 保存済みの最新週 {last_from}"]
    added = 0
    for href, s, e in sorted(links, key=lambda x: x[1]):
        sd = datetime.datetime.strptime(s, "%Y%m%d").date(); ed = datetime.datetime.strptime(e, "%Y%m%d").date()
        mon = monday(sd).isoformat()
        res = requests.get(BASE + href, headers=UA, timeout=60); res.raise_for_status()
        vals, seen = parse(res.content)
        cmp = ""
        for w in doc["weeks"]:
            if w["from"] == mon:
                cmp = " / 保存値との差: " + ", ".join(f"{k}={vals[k] - w[k]:+d}" for k in COLS)
        check.append(f"{s}〜{e}: {vals} 市場={seen}{cmp}")
        if mon > last_from:
            doc["weeks"].append({"from": mon, "to": monday(sd).__add__(datetime.timedelta(days=4)).isoformat() if ed < monday(sd) else ed.isoformat(), **vals})
            last_from = mon; added += 1
    doc["weeks"].sort(key=lambda w: w["from"])
    if added:
        doc["updated"] = datetime.date.today().isoformat()
        doc["note"] = "2026/9/29掲載分以降はJPX公式xlsxの3市場の差引をscripts/fetch_investor_flows.pyが毎週追記。それ以前の週は株探の週次集計から転記(最新週の10/5〜9はJPX公式と一致を確認済み)。"
        json.dump(doc, open(JSON_PATH, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    check.append(f"追記 {added}週")
    open(CHECK_PATH, "w", encoding="utf-8").write("\n".join(check) + "\n")
    print("\n".join(check))


if __name__ == "__main__":
    try:
        main()
    except Exception as ex:
        open(CHECK_PATH, "w", encoding="utf-8").write(f"失敗 {datetime.datetime.utcnow().isoformat()}Z: {ex!r}\n")
        print("失敗", repr(ex)); sys.exit(1)
