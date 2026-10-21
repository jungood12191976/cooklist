#!/usr/bin/env python3
"""確認用: JPX 投資部門別売買状況の一覧ページと最新週のExcelを取得し、構造を data/investor_flows_probe.txt に書く。"""
import re, io, sys, requests
from openpyxl import load_workbook
BASE = "https://www.jpx.co.jp"
IDX = BASE + "/markets/statistics-equities/investor-type/index.html"
UA = {"User-Agent": "Mozilla/5.0 (king-the-quest personal app)"}
out = []
try:
    r = requests.get(IDX, headers=UA, timeout=60); out.append(f"index HTTP {r.status_code} len={len(r.text)}")
    links = re.findall(r'href="([^"]*stock_1_w_\d{8}_\d{8}\.xlsx)"', r.text)
    out.append("xlsx links: " + repr(links[:6]))
    old = re.findall(r'href="([^"]*stock_val_1_\d{6}\.xls)"', r.text)
    out.append("old val links: " + repr(old[:3]))
    if links:
        u = links[0] if links[0].startswith("http") else BASE + links[0]
        x = requests.get(u, headers=UA, timeout=60); out.append(f"xlsx {u} HTTP {x.status_code} bytes={len(x.content)}")
        wb = load_workbook(io.BytesIO(x.content), data_only=True)
        for ws in wb.worksheets:
            out.append(f"=== sheet {ws.title!r} dims={ws.dimensions}")
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 70)):
                cells = [(c.coordinate, c.value) for c in row if c.value is not None]
                if cells: out.append(repr(cells)[:400])
except Exception as e:
    out.append("ERROR " + repr(e))
open("data/investor_flows_probe.txt", "w", encoding="utf-8").write("\n".join(out))
print("\n".join(out)[:3000])
